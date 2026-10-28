"""End-to-end tests for ARO Publication Navigator (Python Playwright; no npm packages needed).

    python3 tests/e2e.py http://localhost:4173                 # desktop (1280×860)
    python3 tests/e2e.py http://localhost:4173 "iPhone 13"     # any Playwright device: touch, mobile viewport

Covers every route, progress persistence and cross-view sync, search, the knowledge graph (select, path,
list mode), handbook modules, the .docx checker (with generated good and bad manuscripts), exports,
keyboard use, mobile layout, dark theme and reduced motion. Exits non-zero on any failure."""
import sys, tempfile, pathlib, traceback
from playwright.sync_api import sync_playwright, expect

BASE = (sys.argv[1] if len(sys.argv) > 1 else 'http://localhost:4173').rstrip('/')
DEVICE = sys.argv[2] if len(sys.argv) > 2 else ''
results: list[tuple[str, bool, str]] = []
IGNORED = ('ERR_TUNNEL_CONNECTION_FAILED', 'fonts.googleapis', 'ERR_NAME_NOT_RESOLVED', 'ERR_INTERNET_DISCONNECTED')


def test(name):
    def deco(fn):
        def run(page):
            try:
                fn(page); results.append((name, True, ''))
            except Exception as e:  # noqa: BLE001
                results.append((name, False, f'{type(e).__name__}: {e}'.splitlines()[0][:300]))
        run.__name__ = name
        TESTS.append(run)
        return fn
    return deco


TESTS = []


def go(page, route):
    page.goto(f'{BASE}/{route}')
    page.wait_for_selector('main#main > *')


def make_docx(path: pathlib.Path, good: bool):
    import docx
    from docx.shared import Pt, Mm
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(215.9), Mm(279.4)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Mm(20))
    st = d.styles['Normal']
    st.font.name = 'Times New Roman' if good else 'Calibri'
    st.font.size = Pt(12 if good else 10)
    st.paragraph_format.line_spacing = 2.0 if good else 1.0
    d.add_heading('Comparing subword tokenizers for an example language', 0)
    d.add_heading('Abstract', 1)
    words = ('This illustrative abstract describes a study of tokenizers for an example language and reports what was done why it matters and what was found. ' * (12 if good else 2)).split()
    d.add_paragraph(' '.join(words[:200] if good else words[:40]) + ('' if good else ' (Smith, 2020)'))
    d.add_paragraph('Index Terms—Byte pair encoding, Kurdish, Language models, Tokenization, Unigram')
    body = ' '.join(['Example sentence about corpora and evaluation of segmentation quality.'] * 40)
    for h in ['Introduction', 'Materials and Methods', 'Results and Discussion', 'Conclusions', 'Acknowledgment']:
        d.add_heading(h, 1)
        for _ in range(5 if good else 1):
            d.add_paragraph(body)
    d.add_heading('References', 1)
    n = 16 if good else 3
    for i in range(n):
        d.add_paragraph(f'Author{chr(65 + i)}, A.B., 20{10 + i}. Example title {i}. Journal of Examples, {i + 1}(2), pp.1-10. Retrieved from http://dx.doi.org/10.0000/ex.{i}.')
    if not good:
        d.add_paragraph('Preprintauthor, X., 2024. A preprint. arXiv preprint arXiv:2401.00001.')
        d.add_paragraph('Contact me at someone@example.org; code at github.com/someone/repo; in our previous work we showed this.')
    d.core_properties.author = '' if good else 'Real Author Name'
    d.save(path)


# ── tests ────────────────────────────────────────────────────────────────────
@test('every primary route renders a heading')
def _(page):
    for r in ['#/', '#/explore', '#/handbook', '#/research', '#/resources', '#/readiness', '#/directory', '#/sources', '#/questions',
              '#/stage/s01', '#/stage/s10', '#/handbook/references', '#/handbook/submission', '#/req/r-abstract', '#/questions/c-anonymization']:
        go(page, r)
        expect(page.locator('main h1').first).to_be_visible()


@test('all ten stages and all handbook chapters open')
def _(page):
    for i in range(1, 11):
        go(page, f'#/stage/s{i:02d}')
        expect(page.locator('.check-row').first).to_be_visible()
    go(page, '#/handbook')
    hrefs = page.eval_on_selector_all('.hb-topic h3 a', 'els => els.map(e => e.getAttribute("href"))')
    assert len(hrefs) >= 20, hrefs
    for h in hrefs:
        go(page, h)
        expect(page.locator('main h1')).to_be_visible()


@test('unknown stage shows a not-found message')
def _(page):
    go(page, '#/stage/nope')
    expect(page.get_by_text('Stage not found')).to_be_visible()


@test('ticking an item persists after reload and updates journey, readiness and graph')
def _(page):
    go(page, '#/stage/s01')
    page.locator('#chk-k-read-guide').check()
    page.wait_for_timeout(200)
    page.reload(); page.wait_for_selector('#chk-k-read-guide')
    expect(page.locator('#chk-k-read-guide')).to_be_checked()
    go(page, '#/')
    expect(page.locator('.dash-tile--green .kpi__value')).to_have_text('1')
    go(page, '#/readiness?show=done')
    expect(page.locator('.ready-group .check-row')).to_have_count(1)
    go(page, '#/explore')
    page.locator('[role=tab]', has_text='Graph').click()
    expect(page.locator('g.gnode[data-id="s01"] .gnode__prog')).to_have_text('1/3')


@test('blocked items warn instead of locking; dependencies unblock')
def _(page):
    go(page, '#/stage/s02')
    row = page.locator('.check-row:has(#chk-k-contribution)')
    expect(row).to_have_attribute('data-status', 'blocked')
    page.locator('#chk-k-experiments-done').check(); page.wait_for_timeout(150)
    expect(page.locator('.check-row:has(#chk-k-contribution)')).to_have_attribute('data-status', 'open')


@test('recording the office answer unblocks awaiting items')
def _(page):
    go(page, '#/stage/s04')
    expect(page.locator('.check-row:has(#chk-k-discipline)')).to_have_attribute('data-status', 'awaiting')
    go(page, '#/questions/c-discipline')
    page.fill('#q-note', 'Office said: tick Computer Science.')
    page.get_by_role('button', name='Save answer').click()
    go(page, '#/stage/s04')
    expect(page.locator('.check-row:has(#chk-k-discipline)')).to_have_attribute('data-status', 'open')


@test('continue where I left off returns to the last route')
def _(page):
    go(page, '#/handbook/abstract')
    go(page, '#/')
    page.locator('[data-continue]').click()
    page.wait_for_url('**/#/handbook/abstract')


@test('global search finds requirements, people and documents; Enter opens')
def _(page):
    go(page, '#/')
    page.keyboard.press('/')
    page.keyboard.type('orcid')
    expect(page.locator('.search-results li').first).to_be_visible()
    page.keyboard.press('Escape')
    page.locator('[data-search-btn]').click()
    page.keyboard.type('Yahya')
    expect(page.locator('.search-results')).to_contain_text('Salah I. Yahya')
    page.locator('.search-input').fill('stage-one template')
    expect(page.locator('.search-results')).to_contain_text('Template')
    page.keyboard.press('Enter')
    expect(page.locator('main h1')).to_be_visible()


@test('graph: select node, shortest path, dependency highlight, list mode')
def _(page):
    if not DEVICE: page.set_viewport_size({'width': 1440, 'height': 900})
    go(page, '#/explore?focus=r-double-blind')
    page.locator('[role=tab]', has_text='Graph').click()
    expect(page.locator('.explore-panel h2')).to_have_text('Double-blind peer review')
    page.get_by_role('button', name='Path from Stage 01').click()
    expect(page.locator('.path-steps li').first).to_have_text('01 Understand ARO')
    assert page.locator('g.gnode.is-lit').count() >= 2
    page.get_by_role('button', name='Everything it needs').click()
    expect(page.locator('.explore-panel')).to_contain_text('prerequisites')
    page.get_by_role('button', name='Zoom in').click(); page.get_by_role('button', name='Fit to view').click()
    page.locator('[role=tab]', has_text='List').click()
    expect(page.locator('.graph-list__group').first).to_be_visible()
    page.locator('.graph-list__node').first.click()
    expect(page.locator('.explore-panel h2')).to_be_visible()


@test('graph relations are logically typed')
def _(page):
    go(page, '#/explore')
    page.locator('[role=tab]', has_text='List').click()
    go(page, '#/explore?focus=res-stage-one')
    expect(page.locator('.explore-panel')).to_contain_text('downloaded from')
    go(page, '#/explore?focus=c-anonymization')
    expect(page.locator('.explore-panel')).to_contain_text('conflicts with')


@test('handbook filter and Harvard practice persist')
def _(page):
    go(page, '#/handbook')
    total = page.locator('.hb-topic').count()
    page.fill('.hb-filter', 'ORCID')
    page.wait_for_timeout(100)
    assert 0 < page.locator('.hb-topic').count() < total
    expect(page.locator('.hb-topics')).to_contain_text('Author information')
    go(page, '#/handbook/references')
    page.locator('#module summary', has_text='Practice').click()
    page.locator('#qc-hv1-b').click(); page.locator('#qchk-hv1').click()
    expect(page.locator('#quiz-hv1 .quiz-fb__bar--correct')).to_be_visible()
    page.reload(); page.wait_for_selector('#quiz-hv1', state='attached')
    expect(page.locator('#qc-hv1-b')).to_have_attribute('data-state', 'correct')


@test('length calculator flags overlong figure counts')
def _(page):
    go(page, '#/handbook/length')
    page.locator('input[aria-label="Words (number)"]').fill('5200')
    page.locator('input[aria-label="Figures (number)"]').fill('9')
    expect(page.locator('#module')).to_contain_text('makes the paper overlong')
    page.locator('input[aria-label="Figures (number)"]').fill('7')
    expect(page.locator('#module')).to_contain_text('7 is fine')


@test('abstract workbench counts words and flags citations')
def _(page):
    go(page, '#/handbook/abstract')
    page.fill('#abs-ta', 'Tokenizers matter (Smith, 2020). ' * 3)
    expect(page.locator('#module')).to_contain_text('citation')
    expect(page.locator('.abs-meter')).to_contain_text('12 words')


@test('structure, AI and anonymisation modules give feedback')
def _(page):
    go(page, '#/handbook/organization')
    page.get_by_role('button', name='Show me').click()
    expect(page.locator('#module')).to_contain_text('Correct!')
    go(page, '#/handbook/ai')
    page.locator('.ai-row').first.get_by_role('button', name='Allowed', exact=True).click()
    expect(page.locator('.ai-row').first).to_contain_text('Correct!')
    go(page, '#/handbook/blind')
    page.get_by_role('button', name='Show me').click()
    expect(page.locator('#module')).to_contain_text('4 of 4 found')


@test('submission walkthrough and review simulator step through')
def _(page):
    go(page, '#/handbook/submission')
    for _ in range(7):
        page.get_by_role('button', name='Continue').click()
    expect(page.locator('#module')).to_contain_text('Step 8 of 8')
    go(page, '#/handbook/revisions')
    page.get_by_role('button', name='Editorial screening').first.click()
    expect(page.locator('.flow-card')).to_contain_text('15%')


@test('resources: filters, real links, downloads labelled honestly')
def _(page):
    go(page, '#/resources')
    page.select_option('select[aria-label="Type"]', 'download')
    cards = page.locator('.res-card')
    assert cards.count() >= 2
    hrefs = page.eval_on_selector_all('.res-card .res-actions a', 'els => els.map(e => e.href)')
    assert all(h.startswith('https://docs.google.com') or h.startswith('https://drive.google.com') or h.startswith('http://publicationethics.org') for h in hrefs), hrefs
    page.select_option('select[aria-label="Type"]', 'page')
    labels = page.eval_on_selector_all('.res-card .res-actions a', 'els => els.map(e => e.textContent)')
    assert all('download' not in l.lower() for l in labels), labels
    with page.expect_download() as dl:
        page.get_by_role('button', name='Download my preparation checklist').click()
    assert dl.value.suggested_filename.endswith('.md')


@test('every requirement cites an official https source')
def _(page):
    go(page, '#/sources')
    urls = page.eval_on_selector_all('.src-table tbody a', 'els => els.map(e => e.href)')
    assert len(urls) >= 40
    assert all(u.startswith('http') for u in urls)
    assert sum('aro.koyauniversity.org' in u for u in urls) >= 20


@test('docx checker: compliant draft mostly passes, non-compliant draft is flagged')
def _(page):
    tmp = pathlib.Path(tempfile.mkdtemp())
    good, bad = tmp / 'good.docx', tmp / 'bad.docx'
    make_docx(good, True); make_docx(bad, False)
    go(page, '#/readiness')
    page.set_input_files('#docx-in', str(good))
    expect(page.locator('.docx-table')).to_be_visible(timeout=10000)
    fails = page.locator('.docx-table .verdict--fail').count()
    assert fails <= 1, f'{fails} fails on the compliant file: ' + page.locator('.docx-table').inner_text()[:600]
    page.set_input_files('#docx-in', str(bad))
    expect(page.locator('.docx-table')).to_contain_text('bad.docx' if False else 'Identity', timeout=10000)
    txt = page.locator('.docx-table').inner_text()
    assert page.locator('.docx-table .verdict--fail, .docx-table .verdict--warn').count() >= 5, txt[:800]
    assert 'Real Author Name' in txt
    page.get_by_role('button', name='Link verdicts to checklist').click()
    expect(page.locator('.auto-chip').first).to_be_visible()


@test('docx checker rejects non-docx files safely')
def _(page):
    tmp = pathlib.Path(tempfile.mkdtemp()) / 'notes.docx'
    tmp.write_text('not a zip')
    go(page, '#/readiness')
    page.set_input_files('#docx-in', str(tmp))
    expect(page.get_by_role('alert')).to_contain_text('Could not check')


@test('readiness exports Markdown and progress backup')
def _(page):
    go(page, '#/readiness')
    with page.expect_download() as dl:
        page.get_by_role('button', name='Markdown').click()
    p = dl.value.path(); text = pathlib.Path(p).read_text()
    assert '# ARO submission preparation checklist' in text and '[x] Read the Guide to Authors' in text
    with page.expect_download() as dl2:
        page.get_by_role('button', name='Back up progress').click()
    assert '"checked"' in pathlib.Path(dl2.value.path()).read_text()


@test('directory search and mailto draft without personal emails')
def _(page):
    go(page, '#/directory')
    page.fill('input[aria-label="Search the editorial directory"]', 'Portugal')
    expect(page.locator('.person')).to_have_count(1)
    href = page.get_by_role('link', name='Open in my email app').get_attribute('href')
    assert href.startswith('mailto:aro.journal@koyauniversity.org?subject=')
    mails = page.eval_on_selector_all('a[href^="mailto:"]', 'els => [...new Set(els.map(e => e.href.split("?")[0]))]')
    assert mails == ['mailto:aro.journal@koyauniversity.org'], mails


@test('research criteria tick and persist')
def _(page):
    go(page, '#/research?focus=sec-methods')
    box = page.locator('#rc-sec-methods-0')
    box.check(); page.reload(); page.wait_for_selector('#rc-sec-methods-0')
    expect(page.locator('#rc-sec-methods-0')).to_be_checked()


@test('guided tour walks all seven areas')
def _(page):
    go(page, '#/')
    page.get_by_role('button', name='Take the tour').click()
    for _ in range(6):
        page.locator('.tour-modal').get_by_role('button', name='Continue').click()
    expect(page.locator('.tour-modal')).to_contain_text('Step 7 of 7')
    page.locator('.tour-modal').get_by_role('button', name='Finish').click()
    expect(page.locator('.tour-modal')).to_have_count(0)


@test('keyboard: skip link, journey popover, Enter on graph node')
def _(page):
    go(page, '#/')
    page.keyboard.press('Tab')
    expect(page.locator('.skip-link')).to_be_focused()
    page.locator('.path__item[aria-controls="pop-s01"]').focus(); page.keyboard.press('Enter')
    expect(page.locator('#pop-s01')).to_be_visible()
    go(page, '#/explore')
    page.locator('[role=tab]', has_text='Graph').click()
    page.locator('g.gnode[data-id="s03"]').focus(); page.keyboard.press('Enter')
    expect(page.locator('.explore-panel h2')).to_contain_text('Structure your paper')


@test('GitHub link: logo and username in the top bar, opens github.com/SawabS')
def _(page):
    go(page, '#/')
    a = page.locator('.app-bar a.gh-link')
    expect(a).to_be_visible()
    assert a.get_attribute('href') == 'https://github.com/SawabS' and a.get_attribute('target') == '_blank'
    expect(a.locator('svg use')).to_have_attribute('href', '#a-github')
    name = a.locator('.gh-link__name')
    if DEVICE: expect(name).to_be_hidden(); bb = a.bounding_box(); assert bb['width'] >= 40 and bb['height'] >= 40, bb
    else: expect(name).to_have_text('SawabS')


@test('theme toggle switches and persists')
def _(page):
    go(page, '#/')
    page.get_by_role('button', name='Switch light or dark theme').click()
    t1 = page.evaluate('document.documentElement.dataset.theme')
    page.reload(); page.wait_for_selector('main#main > *')
    assert page.evaluate('document.documentElement.dataset.theme') == t1


@test('reset progress clears everything')
def _(page):
    go(page, '#/')
    page.once('dialog', lambda d: d.accept())
    page.get_by_role('button', name='Reset progress').click()
    page.wait_for_timeout(200)
    expect(page.locator('.dash-tile--green .kpi__value')).to_have_text('0')


def phone_tests(page):
    """Touch-specific checks, run only with a device argument."""
    def t(name, fn):
        try: fn(); results.append((name, True, ''))
        except Exception as e: results.append((name, False, f'{type(e).__name__}: ' + ' | '.join(l.strip() for l in str(e).splitlines() if l.strip())[:700]))  # noqa: E701

    def tabbar():
        go(page, '#/')
        expect(page.locator('.tabbar')).to_be_visible()
        expect(page.locator('.main-nav')).to_be_hidden()
        for label, url in [('Graph', '#/explore'), ('Rules', '#/handbook'), ('Paper', '#/research'), ('Files', '#/resources'), ('Ready', '#/readiness'), ('ARO', '#/directory'), ('Journey', '#/')]:
            page.locator('.tabbar a', has_text=label).tap()
            page.wait_for_url('**/' + url)
            expect(page.locator(f'.tabbar a[href="{url}"]')).to_have_attribute('aria-current', '')
        # the fixed bar never hides the end of the page
        go(page, '#/stage/s01')
        page.evaluate('window.scrollTo(0, document.body.scrollHeight)'); page.wait_for_timeout(200)
        bar = page.locator('.tabbar').bounding_box(); foot = page.locator('.app-footer').bounding_box()
        assert foot['y'] + foot['height'] <= bar['y'] + 1, (foot, bar)
    t('phone: bottom tab bar navigates every area and never covers content', tabbar)

    def graph_touch():
        go(page, '#/explore')
        expect(page.locator('.graph-list')).to_be_visible()          # list view by default on phones
        expect(page.locator('.explore-panel')).to_be_hidden()        # no empty panel taking space
        page.locator('.graph-list__node', has_text='01 Understand ARO').tap()
        sheet = page.locator('.explore-panel.has-sel')
        expect(sheet).to_be_visible()
        box = sheet.bounding_box(); vh = page.viewport_size['height']
        assert box['y'] > vh * 0.2 and box['y'] + box['height'] <= vh + 1, box   # a bottom sheet, fully on screen
        y = page.evaluate('window.scrollY')
        sheet.get_by_role('button', name='Clear selection').tap()
        expect(page.locator('.explore-panel')).to_be_hidden()
        assert abs(page.evaluate('window.scrollY') - y) < 5               # selecting did not throw the page around
        page.locator('[role=tab]', has_text='Graph').tap()
        expect(page.locator('svg.graph-svg')).to_be_visible()
        node = page.locator('g.gnode[data-id="s01"]')
        node.tap()
        expect(page.locator('.explore-panel h2')).to_contain_text('Understand ARO')
        page.locator('.explore-panel').get_by_role('button', name='Clear selection').tap()
        expect(page.locator('.explore-panel')).to_be_hidden()
        k0 = page.evaluate("+document.querySelector('.graph-world').getAttribute('transform').match(/scale\\(([\\d.]+)/)[1]")
        page.locator('svg.graph-svg').scroll_into_view_if_needed()
        page.evaluate("document.querySelector('svg.graph-svg').scrollIntoView({block: 'center'})"); page.wait_for_timeout(200)
        box = page.locator('svg.graph-svg').bounding_box()
        cx, cy = box['x'] + box['width'] / 2, box['y'] + box['height'] / 2
        cdp = page.context.new_cdp_session(page)
        def touch(kind, d):
            pts = [] if kind == 'touchEnd' else [{'x': cx - d, 'y': cy, 'id': 1}, {'x': cx + d, 'y': cy, 'id': 2}]
            cdp.send('Input.dispatchTouchEvent', {'type': kind, 'touchPoints': pts})
        touch('touchStart', 30)
        for d in range(40, 130, 10): touch('touchMove', d)
        touch('touchEnd', 0)
        page.wait_for_timeout(100)
        k1 = page.evaluate("+document.querySelector('.graph-world').getAttribute('transform').match(/scale\\(([\\d.]+)/)[1]")
        assert k1 > k0 * 2, (k0, k1)                                       # pinch-out zoomed in
    t('phone: graph list view, bottom-sheet details, tap on nodes, pinch zoom', graph_touch)

    def targets():
        small = []
        for r in ['#/', '#/stage/s04', '#/explore', '#/handbook/references', '#/handbook/length', '#/resources', '#/readiness', '#/directory']:
            go(page, r)
            small += page.evaluate(f"""() => [...document.querySelectorAll('main button, main select, main input:not([type=checkbox]):not([type=file]):not(.visually-hidden), main .btn, main .btn3d, main .chip, .tabbar a, .app-bar button')]
              .filter(e => {{ const b = e.getBoundingClientRect(); return b.width > 0 && b.height > 0 && (b.height < 36 || b.width < 36) && !e.closest('details:not([open]) > :not(summary)') && !e.classList.contains('link-btn'); }})
              .map(e => '{r} ' + (e.getAttribute('aria-label') || e.textContent.trim().slice(0, 30)) + ' ' + Math.round(e.getBoundingClientRect().width) + 'x' + Math.round(e.getBoundingClientRect().height))""")
        assert not small, small[:8]
    t('phone: every button, field and tab is at least 36 px in both directions', targets)

    def readable():
        go(page, '#/handbook/references')
        size = page.evaluate("parseFloat(getComputedStyle(document.querySelector('.req-rule')).fontSize)")
        assert size >= 15, size
        small = page.evaluate("""() => [...document.querySelectorAll('main p, main li, main td, main dd, main label')].filter(e => e.offsetParent && e.textContent.trim() && parseFloat(getComputedStyle(e).fontSize) < 12).length""")
        assert small == 0, small
    t('phone: body text is 12 px or larger', readable)

    def forms():
        go(page, '#/stage/s01')
        page.locator('label[for="chk-k-read-example"]').tap()
        expect(page.locator('#chk-k-read-example')).to_be_checked()
        go(page, '#/handbook/length')
        page.locator('input[aria-label="Figures (number)"]').fill('20')
        expect(page.locator('#module')).to_contain_text('overlong')
        go(page, '#/handbook/references')
        page.locator('.jump-module').tap()
        page.wait_for_timeout(300)
        assert page.locator('#module').bounding_box()['y'] < page.viewport_size['height']
    t('phone: tapping labels, typing in fields and the jump link all work', forms)


def mobile_and_motion(browser):
    ctx = browser.new_context(viewport={'width': 320, 'height': 700}, reduced_motion='reduce', color_scheme='dark')
    page = ctx.new_page()
    try:
        for r in ['#/', '#/stage/s04', '#/explore', '#/handbook/references', '#/research', '#/resources', '#/readiness', '#/directory', '#/questions/c-anonymization']:
            go(page, r)
            ov = page.evaluate('document.documentElement.scrollWidth - document.documentElement.clientWidth')
            assert ov <= 1, f'{r}: {ov}px horizontal overflow at 320px'
        go(page, '#/explore')
        expect(page.locator('.graph-list')).to_be_visible()  # list/tree alternative by default on phones
        expect(page.locator('.tabbar')).to_be_visible()
        dur = page.evaluate("getComputedStyle(document.querySelector('main')).animationDuration")
        assert dur in ('1e-06s', '0.000001s', '0s') or float(dur.rstrip('s')) < 0.01, dur
        results.append(('320px dark reduced-motion: no overflow, list graph, tab bar, motion off', True, ''))
    except Exception as e:  # noqa: BLE001
        results.append(('320px dark reduced-motion: no overflow, list graph, tab bar, motion off', False, str(e)[:300]))
    ctx.close()


with sync_playwright() as p:
    browser = p.chromium.launch()
    if DEVICE:
        dev = dict(p.devices[DEVICE]); dev.pop('default_browser_type', None)
        ctx = browser.new_context(**dev, accept_downloads=True)
    else:
        ctx = browser.new_context(viewport={'width': 1280, 'height': 860}, accept_downloads=True)
    page = ctx.new_page()
    errors: list[str] = []
    page.on('pageerror', lambda e: errors.append(str(e)))
    page.on('console', lambda m: m.type == 'error' and not any(s in m.text for s in IGNORED) and errors.append(m.text))
    go(page, '#/'); page.evaluate('localStorage.clear()'); page.reload()
    for t in TESTS:
        t(page)
    results.append(('no JavaScript errors in the console', not errors, '; '.join(errors[:3])))
    if DEVICE:
        go(page, '#/'); page.evaluate('localStorage.clear()'); page.reload()
        phone_tests(page)
    mobile_and_motion(browser)
    browser.close()

w = max(len(n) for n, _, _ in results)
for n, ok, msg in results:
    print(f"{'PASS' if ok else 'FAIL'}  {n.ljust(w)}  {msg}")
failed = sum(not ok for _, ok, _ in results)
print(f'\n{len(results) - failed}/{len(results)} passed' + (f' on {DEVICE}' if DEVICE else ' on desktop'))
sys.exit(1 if failed else 0)
